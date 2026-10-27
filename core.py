"""Build a PowerPoint from Olympus/Evident confocal *.tif.frames exports."""
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image, ImageDraw

import imaging
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# Label text color per channel position (DAPI, ch2, ch3, ch4) -- brightened so they read on black.
CHANNEL_COLORS = ["4C8DFF", "FF4D4D", "3DDC63", "D0D0D0"]


@dataclass
class Options:
    recursive: bool = True
    show_merge: bool = True
    show_title: bool = True
    show_filename: bool = True
    label_size: int = 12
    image_px: int = 900
    prefix_order: list = field(default_factory=list)  # e.g. ["BRPC-26", "26-T"]
    scale_bar: bool = True


@dataclass
class Sample:
    base: str
    prefix: str
    names: list
    merge: Path
    channels: list
    txt: Path = None
    um_per_px: float = None

    @property
    def key(self):
        return str(self.merge)


@dataclass
class ChannelSetting:
    lo: int = 0
    hi: int = 255
    locked: bool = False
    visible: bool = True

    def is_default(self):
        return self.lo == 0 and self.hi == 255 and self.visible


@dataclass
class SampleSettings:
    channels: list
    include: bool = True
    roi: tuple = None  # (x, y, size) as fractions of image width

    def is_untouched(self):
        return all(c.is_default() for c in self.channels)


def default_settings(samples):
    return {s.key: SampleSettings(channels=[ChannelSetting() for _ in s.names]) for s in samples}


def set_include_all(settings, value):
    for st in settings.values():
        st.include = value


def apply_to_all(settings, source_key):
    """Copy the source sample's display windows to every other sample, skipping locked channels."""
    src = settings[source_key]
    for key, st in settings.items():
        if key == source_key:
            continue
        for ch, sch in zip(st.channels, src.channels):
            if not ch.locked:
                ch.lo, ch.hi = sch.lo, sch.hi


def parse_names(base):
    """Channel names from filename: DAPI + following tokens, up to magnification/frame number."""
    toks = base.split("_")
    if "DAPI" not in toks:
        return toks[:0], base
    i = toks.index("DAPI")
    names = []
    for t in toks[i:]:
        if re.fullmatch(r"\d+XL?", t) or re.fullmatch(r"\d+", t):
            break
        names.append(t)
    return names[:4], "_".join(toks[:i])


def find_samples(folder, recursive=True):
    folder = Path(folder)
    dirs = sorted(folder.rglob("*.tif.frames") if recursive else folder.glob("*.tif.frames"))
    samples, skipped = [], []
    for d in dirs:
        if not d.is_dir():
            continue
        base = d.name[: -len(".tif.frames")]
        merge = d / f"{base}_T001.tif"
        chans = [d / f"{base}_C00{k}T001.tif" for k in range(1, 5)]
        names, prefix = parse_names(base)
        if not merge.exists() or not names:
            skipped.append((base, "no merged image (T001) or no DAPI in name"))
            continue
        chans = chans[: len(names)]
        if not all(c.exists() for c in chans):
            skipped.append((base, "missing channel image"))
            continue
        txt = d / f"{base}.txt"
        samples.append(Sample(base, prefix, names, merge, chans, txt if txt.exists() else None,
                              imaging.parse_um_per_pixel(txt)))
    return samples, skipped


def order_samples(samples, prefix_order):
    def key(s):
        for i, p in enumerate(prefix_order):
            if s.base.lower().startswith(p.lower()):
                return (i, s.base)
        return (len(prefix_order), s.base)
    return sorted(samples, key=key)


def _load(path, px):
    return Image.open(path).convert("RGB").resize((px, px), Image.LANCZOS)


def _jpeg(path, px):
    return _jpeg_img(_load(path, px), px)


def _jpeg_img(im, px):
    im = im.convert("RGB").resize((px, px), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=88)
    buf.seek(0)
    return buf


def _text(slide, x, y, w, h, text, size, bold=False, color="1F2937", name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = tf.paragraphs[0].add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold = "Calibri", Pt(size), bold
    r.font.color.rgb = RGBColor.from_string(color)
    if name:
        tb.name = name


def _label(slide, pic, text, color, size):
    """Black box with colored text in the top-right corner of a picture."""
    pad = 0.06
    w = max(0.45, len(text) * size * 0.0095 + 2 * pad)  # rough Calibri bold width, inches
    h = size / 72 * 1.25 + 2 * pad * 0.6
    inset = 0.06
    right = (pic.left + pic.width) / 914400
    tb = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(right - inset - w), Inches(pic.top / 914400 + inset), Inches(w), Inches(h))
    tb.name = f"Label {text}"
    tb.fill.solid()
    tb.fill.fore_color.rgb = RGBColor(0, 0, 0)
    tb.line.fill.background()
    tb.shadow.inherit = False
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(pad * 0.6)
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold = "Calibri", Pt(size), True
    r.font.color.rgb = RGBColor.from_string(color)


def _merge_image(s, st, px, opts):
    """Original merge when nothing changed; otherwise rebuilt from the windowed visible channels."""
    visible = [(c, p) for c, p in zip(st.channels, s.channels) if c.visible]
    if st.is_untouched() or not visible:
        im = _load(s.merge, px)
    else:
        im = imaging.composite([imaging.apply_window(_load(p, px), c.lo, c.hi) for c, p in visible])
    if opts.scale_bar and s.um_per_px:
        # bar length is for the full-resolution pixel size, so scale um/px to the preview size
        full = Image.open(s.merge).width
        im = imaging.draw_scale_bar(im, s.um_per_px * full / px)
    return im


def build_deck(samples, out_path, opts, progress=lambda i, n, msg: None, settings=None):
    settings = settings or default_settings(samples)
    samples = [s for s in samples if settings[s.key].include]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    blank = prs.slide_layouts[6]
    n = len(samples)
    for idx, s in enumerate(samples, 1):
        progress(idx - 1, n, s.base)
        st = settings[s.key]
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        if opts.show_title:
            _text(slide, 0.5, 0.25, 9, 0.45, s.prefix.replace("_", " ") or s.base, 24, True, name="Title")
        if opts.show_filename:
            _text(slide, 0.5, 0.72, 9, 0.25, s.base, 10, False, "6B7280", name="Filename")
        top, size, gap = 1.05, 4.2, 0.1
        panel = (size - gap) / 2
        if opts.show_merge:
            left_merge, left_grid = 0.65, 0.65 + size + 0.3
            pic = slide.shapes.add_picture(_jpeg_img(_merge_image(s, st, opts.image_px, opts), opts.image_px),
                                           Inches(left_merge), Inches(top), Inches(size), Inches(size))
            pic.name = "Merge"
            if st.roi:
                rx, ry, rs = imaging.clamp_roi(st.roi)
                box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left_merge + rx * size), Inches(top + ry * size),
                                             Inches(rs * size), Inches(rs * size))
                box.name = "ROI"
                box.fill.background()
                box.line.color.rgb = RGBColor(255, 200, 0)
                box.line.width = Pt(1.5)
                box.shadow.inherit = False
        else:
            left_grid = (10 - size) / 2
        shown = [(k, s.names[k], s.channels[k], st.channels[k]) for k in range(len(s.names)) if st.channels[k].visible]
        for j, (k, name, path, cs) in enumerate(shown):
            x = left_grid + (j % 2) * (panel + gap)
            y = top + (j // 2) * (panel + gap)
            im = imaging.crop_roi(imaging.apply_window(Image.open(path), cs.lo, cs.hi), st.roi)
            pic = slide.shapes.add_picture(_jpeg_img(im, opts.image_px), Inches(x), Inches(y), Inches(panel), Inches(panel))
            pic.name = name
            _label(slide, pic, name, CHANNEL_COLORS[k], opts.label_size)
    prs.save(out_path)
    progress(n, n, "done")


FIG_DPI = 300  # the figure uses the slide layout (10 x 5.625 in) rendered at this resolution


def render_figure(s, st, opts):
    """One PNG-sized image laid out like the slide: title, name, merge (scale bar + ROI box), labelled channel panels."""
    d = FIG_DPI
    inch = lambda v: round(v * d)
    fig = Image.new("RGB", (inch(10), inch(5.625)), "white")
    draw = ImageDraw.Draw(fig)
    if opts.show_title:
        draw.text((inch(0.5), inch(0.25) + inch(0.45) // 2), s.prefix.replace("_", " ") or s.base, fill="#1F2937",
                  font=imaging.load_font(round(24 / 72 * d), True), anchor="lm")
    if opts.show_filename:
        draw.text((inch(0.5), inch(0.72) + inch(0.25) // 2), s.base, fill="#6B7280",
                  font=imaging.load_font(round(10 / 72 * d)), anchor="lm")
    top, size, gap = 1.05, 4.2, 0.1
    panel = (size - gap) / 2
    if opts.show_merge:
        left_grid = 0.65 + size + 0.3
        mpx = inch(size)
        merge = _merge_image(s, st, mpx, opts)
        if st.roi:
            x, y, side = imaging.clamp_roi(st.roi)
            ImageDraw.Draw(merge).rectangle((x * mpx, y * mpx, (x + side) * mpx, (y + side) * mpx),
                                            outline=(255, 200, 0), width=max(2, d // 100))
        fig.paste(merge, (inch(0.65), inch(top)))
    else:
        left_grid = (10 - size) / 2
    font = imaging.load_font(round(opts.label_size / 72 * d), True)
    shown = [k for k in range(len(s.names)) if st.channels[k].visible]
    for j, k in enumerate(shown):
        cs = st.channels[k]
        x0, y0 = inch(left_grid + (j % 2) * (panel + gap)), inch(top + (j // 2) * (panel + gap))
        im = imaging.crop_roi(imaging.apply_window(Image.open(s.channels[k]), cs.lo, cs.hi), st.roi)
        pw = inch(panel)
        fig.paste(im.convert("RGB").resize((pw, pw), Image.LANCZOS), (x0, y0))
        l, t, r, b = draw.textbbox((0, 0), s.names[k], font=font)
        padx, pady, inset = inch(0.06), inch(0.04), inch(0.06)
        bw, bh = r - l + 2 * padx, b - t + 2 * pady
        bx, by = x0 + pw - inset - bw, y0 + inset
        draw.rectangle((bx, by, bx + bw, by + bh), fill="black")
        draw.text((bx + bw // 2, by + bh // 2), s.names[k], fill="#" + CHANNEL_COLORS[k], font=font, anchor="mm")
    return fig


def export_pngs(samples, out_dir, opts, progress=lambda i, n, msg: None, settings=None):
    """Write one unified figure PNG per included sample: <base>.png."""
    settings = settings or default_settings(samples)
    samples = [s for s in samples if settings[s.key].include]
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    n = len(samples)
    for idx, s in enumerate(samples, 1):
        progress(idx - 1, n, s.base)
        render_figure(s, settings[s.key], opts).save(out_dir / f"{s.base}.png")
    progress(n, n, "done")
    return n
